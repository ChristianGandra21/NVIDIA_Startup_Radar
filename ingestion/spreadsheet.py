from __future__ import annotations
from datetime import datetime
from io import BytesIO

import httpx
import openpyxl

from services.scraping.schema import StartupProfile, SourceEvidence

INOVATIVA_URL = "https://docs.google.com/spreadsheets/d/12c_-CAkzIbkVs7vUeXUyEsbhNMrx__QZpM4hivS0WP0/pub?output=xlsx"
SHEET_NAME = "Startups"

COL_ANO = 0
COL_CICLO = 1
COL_PROGRAMA = 2
COL_NOME = 3
COL_SITE = 4
COL_ESTADO = 5
COL_STATUS = 6
COL_AREA = 7


def download_xlsx(url: str = INOVATIVA_URL) -> bytes:
    resp = httpx.get(url, follow_redirects=True)
    resp.raise_for_status()
    return resp.content


def parse_xlsx(content: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    ws = wb[SHEET_NAME]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(
            {
                "ano": row[COL_ANO],
                "ciclo": row[COL_CICLO],
                "programa": row[COL_PROGRAMA],
                "nome": row[COL_NOME],
                "site": row[COL_SITE],
                "estado": row[COL_ESTADO],
                "status": row[COL_STATUS],
                "area": row[COL_AREA],
            }
        )
    return rows


def _clean_str(val: object) -> str:
    if val is None:
        return ""
    if isinstance(val, float):
        if val != val:
            return ""
        return str(val).replace("e+", "e")
    return str(val).strip()


def row_to_profile(row: dict, row_number: int) -> StartupProfile:
    nome = _clean_str(row["nome"])
    site = _clean_str(row["site"])

    ano = row["ano"]
    if ano is not None:
        ano = int(ano)

    ciclo = row["ciclo"]
    ciclo_str = None
    if ciclo is not None and isinstance(ciclo, datetime):
        ciclo_str = ciclo.strftime("%Y-%m-%d")
    elif ciclo is not None:
        ciclo_str = str(ciclo)

    source = SourceEvidence(
        url=INOVATIVA_URL,
        extraction_method="spreadsheet_inovativa",
        raw_excerpt=f"linha {row_number} da planilha InovAtiva",
    )

    return StartupProfile(
        name=nome,
        website=site if site else None,
        sector=row.get("area"),
        business_area=row.get("area"),
        state=row.get("estado"),
        program=row.get("programa"),
        cohort_year=ano,
        cohort_cycle=ciclo_str,
        inovativa_status=row.get("status"),
        sources=[source],
    )


def ingest_inovativa(url: str = INOVATIVA_URL) -> list[StartupProfile]:
    content = download_xlsx(url)
    rows = parse_xlsx(content)
    profiles = []
    for i, row in enumerate(rows, start=2):
        profile = row_to_profile(row, i)
        profiles.append(profile)
    return profiles


def ingest_sample(url: str = INOVATIVA_URL, n: int = 5) -> list[StartupProfile]:
    content = download_xlsx(url)
    wb = openpyxl.load_workbook(BytesIO(content), read_only=True)
    ws = wb[SHEET_NAME]
    profiles = []
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=n + 1, values_only=True), start=2):
        row_dict = {
            "ano": row[COL_ANO],
            "ciclo": row[COL_CICLO],
            "programa": row[COL_PROGRAMA],
            "nome": row[COL_NOME],
            "site": row[COL_SITE],
            "estado": row[COL_ESTADO],
            "status": row[COL_STATUS],
            "area": row[COL_AREA],
        }
        profiles.append(row_to_profile(row_dict, i))
    return profiles
